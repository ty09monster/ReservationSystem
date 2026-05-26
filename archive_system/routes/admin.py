import json
import logging
from io import BytesIO
from datetime import datetime, date, timedelta
from flask import (
    Blueprint, render_template, request, redirect, url_for,
    flash, session, jsonify, send_file
)
from werkzeug.security import generate_password_hash, check_password_hash
from sqlalchemy import or_, case, func
from ..extensions import db
from ..models import (
    Admin, Role, Reservation, User, Announcement, SystemConfig,
    Venue, VenueTimeSlot, Attachment, ArchiveRequest,
    VenueTimeSlotDisabledDate, CancelRequest, ApprovalStaff, Guide,
    SystemLog
)

logger = logging.getLogger(__name__)

admin_bp = Blueprint('admin', __name__, url_prefix='/admin')
MIN_PASSWORD_LENGTH = 8

XIAOSHI_CATEGORIES = ['校史馆', '标本馆']
ARCHIVE_CATEGORIES = ['档案馆']

AVAILABLE_PERMISSIONS = [
    ('reservation', '预约记录'),
    ('users', '预约人员'),
    ('approval_staff', '审批人员'),
    ('guides', '讲解员管理'),
    ('notice', '公告管理'),
    ('venue', '场馆管理'),
    ('account', '账号管理'),
    ('roles', '角色权限'),
    ('archive', '档案申请'),
    ('cancel', '撤销申请'),
    ('stats', '数据统计'),
    ('export', '数据导出'),
]


@admin_bp.before_request
def check_admin_status():
    whitelist = {'admin.login', 'admin.logout', 'static', 'admin.get_cancel_request_detail'}
    if request.endpoint in whitelist:
        return

    if not session.get("admin_logged_in"):
        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({"error": "请先登录"}), 401
        return redirect(url_for('admin.login'))

    admin_id = session.get("admin_id")
    security_token = session.get("security_token")
    current_admin = db.session.get(Admin, admin_id)

    if not current_admin:
        session.clear()
        flash("您的账号已被删除，会话中断")
        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({"error": "账号已被删除"}), 401
        return redirect(url_for('admin.login'))

    if current_admin.password_hash and current_admin.password_hash[-6:] != security_token:
        session.clear()
        flash("密码已变更，请重新登录")
        if request.is_json or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({"error": "密码已变更"}), 401
        return redirect(url_for('admin.login'))


def _get_management_area():
    area = request.args.get('area', '').strip()
    if area in ('xiaoshi', 'archive'):
        session['management_area'] = area
        return area
    area = session.get('management_area', 'xiaoshi')
    if area not in ('xiaoshi', 'archive'):
        area = 'xiaoshi'
    session['management_area'] = area
    return area


def _get_area_categories():
    area = _get_management_area()
    if area == 'archive':
        return ARCHIVE_CATEGORIES
    return XIAOSHI_CATEGORIES


def _get_area_label():
    area = _get_management_area()
    return '档案馆管理区' if area == 'archive' else '校史馆管理区'


def _check_permission(perm_key):
    admin_id = session.get("admin_id")
    admin_obj = db.session.get(Admin, admin_id)
    if not admin_obj:
        return False
    if admin_obj.is_super:
        return True
    if admin_obj.role:
        return admin_obj.role.has_permission(perm_key)
    return False


def _write_log(log_type, module, action, detail='', result='成功'):
    try:
        operator = session.get('admin_username', 'unknown')
        operator_ip = request.remote_addr or ''
        log_entry = SystemLog(
            log_type=log_type,
            operator=operator,
            operator_ip=operator_ip,
            module=module,
            action=action,
            detail=detail[:500] if detail else '',
            result=result
        )
        db.session.add(log_entry)
        db.session.commit()
    except Exception:
        db.session.rollback()
        logger.exception("写入日志失败")


@admin_bp.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username")
        password = request.form.get("password")
        admin = Admin.query.filter_by(username=username).first()

        if admin and check_password_hash(admin.password_hash, password):
            session["admin_logged_in"] = True
            session["admin_id"] = admin.id
            session["admin_username"] = admin.username
            session["is_super"] = admin.is_super
            session["security_token"] = admin.password_hash[-6:]
            if 'management_area' not in session:
                session['management_area'] = 'xiaoshi'
            _write_log('login', '系统登录', '管理员登录', f'{username} 登录成功')
            return redirect(url_for("admin.dashboard"))

        _write_log('login', '系统登录', '登录失败', f'{username} 登录失败：用户名或密码错误', '失败')
        flash("用户名或密码错误")
    return render_template("admin_login.html")


@admin_bp.route("/dashboard")
def dashboard():
    active_tab = request.args.get('active_tab', '').strip()
    area = _get_management_area()
    categories = _get_area_categories()

    announcements = Announcement.query.order_by(Announcement.created_at.desc()).all()
    config = SystemConfig.query.first()
    venues = Venue.query.filter(
        Venue.category.in_(categories)
    ).order_by(Venue.id).all()

    admin_list = []
    roles = []
    if session.get("is_super"):
        admin_list = Admin.query.all()
        roles = Role.query.all()

    approval_staff_list = ApprovalStaff.query.order_by(ApprovalStaff.created_at.desc()).all()
    guide_list = Guide.query.order_by(Guide.created_at.desc()).all()

    cancel_requests = []
    if active_tab == "cancel":
        cancel_page = request.args.get('cancel_page', 1, type=int)
        cancel_query = CancelRequest.query.join(Reservation).join(User).filter(
            Reservation.venue.has(Venue.category.in_(categories))
        ).order_by(
            case((CancelRequest.status == '待处理', 0), else_=1),
            CancelRequest.created_at.desc()
        )
        cancel_pagination = cancel_query.paginate(page=cancel_page, per_page=10, error_out=False)
        cancel_requests = cancel_pagination.items

    blacklisted_users_count = User.query.filter_by(is_blacklisted=True).count()

    pending_archive_count = ArchiveRequest.query.filter(
        ArchiveRequest.status.in_(['待审核', '待处理'])
    ).count()
    pending_cancel_count = CancelRequest.query.join(Reservation).filter(
        CancelRequest.status == '待处理',
        Reservation.venue.has(Venue.category.in_(categories))
    ).count()

    return render_template(
        "admin_dashboard.html",
        announcements=announcements,
        config=config,
        venues=venues,
        admin_list=admin_list,
        roles=roles,
        active_tab=active_tab,
        cancel_requests=cancel_requests if active_tab == "cancel" else [],
        pending_archive_count=pending_archive_count,
        pending_cancel_count=pending_cancel_count,
        blacklisted_users_count=blacklisted_users_count,
        approval_staff_list=approval_staff_list,
        guide_list=guide_list,
        management_area=area,
        area_label=_get_area_label(),
        available_permissions=AVAILABLE_PERMISSIONS,
        xiaoshi_categories=json.dumps(XIAOSHI_CATEGORIES),
        archive_categories=json.dumps(ARCHIVE_CATEGORIES),
    )


@admin_bp.route("/api/reservations")
def api_reservations():
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 15, type=int)
    keyword = request.args.get('keyword', '').strip()
    status_filter = request.args.get('status', '').strip()
    res_type_filter = request.args.get('res_type', '').strip()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    venue_id = request.args.get('venue_id', '').strip()
    visit_start = request.args.get('visit_start', '').strip()
    visit_end = request.args.get('visit_end', '').strip()
    categories = _get_area_categories()

    query = Reservation.query.join(User).join(Venue).filter(
        Venue.category.in_(categories)
    )

    if keyword:
        query = query.filter(
            or_(
                User.name.contains(keyword),
                User.phone.contains(keyword),
                User.id_card.contains(keyword),
                Reservation.group_name.contains(keyword),
                Reservation.visiting_unit.contains(keyword)
            )
        )

    if status_filter:
        query = query.filter(Reservation.status == status_filter)

    if res_type_filter:
        query = query.filter(Reservation.res_type == res_type_filter)

    if start_date:
        query = query.filter(Reservation.visit_date >= start_date)

    if end_date:
        query = query.filter(Reservation.visit_date <= end_date)

    if visit_start:
        query = query.filter(func.date(Reservation.created_at) >= visit_start)

    if visit_end:
        query = query.filter(func.date(Reservation.created_at) <= visit_end)

    if venue_id:
        query = query.filter(Reservation.venue_id == int(venue_id))

    query = query.order_by(Reservation.created_at.desc())

    pagination = query.paginate(page=page, per_page=per_page, error_out=False)

    return jsonify({
        "reservations": [
            {
                "id": r.id,
                "user_name": r.user.name,
                "user_phone": r.user.phone,
                "user_id_card": r.user.id_card,
                "user_id_type": r.user.id_type,
                "res_type": r.res_type,
                "group_name": r.group_name or '',
                "visiting_unit": r.visiting_unit or '',
                "visitor_count": r.visitor_count or 1,
                "venue_name": r.venue.name,
                "venue_category": r.venue.category,
                "visit_date": r.visit_date.strftime('%Y-%m-%d') if r.visit_date else '',
                "visit_time": r.visit_time or '',
                "status": r.status,
                "reject_reason": r.reject_reason or '',
                "guide_info": r.guide_info or '',
                "need_guide": r.need_guide,
                "created_at": r.created_at.strftime('%Y-%m-%d %H:%M'),
                "verified_at": r.verified_at.strftime('%Y-%m-%d %H:%M') if r.verified_at else '',
            }
            for r in pagination.items
        ],
        "pagination": {
            "page": pagination.page,
            "pages": pagination.pages,
            "has_prev": pagination.has_prev,
            "has_next": pagination.has_next,
            "total": pagination.total
        }
    })


@admin_bp.route("/export-reservations")
def export_reservations():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    keyword = request.args.get('keyword', '').strip()
    status_filter = request.args.get('status', '').strip()
    res_type_filter = request.args.get('res_type', '').strip()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    venue_id = request.args.get('venue_id', '').strip()
    visit_start = request.args.get('visit_start', '').strip()
    visit_end = request.args.get('visit_end', '').strip()
    categories = _get_area_categories()

    query = Reservation.query.join(User).join(Venue).filter(
        Venue.category.in_(categories)
    )

    if keyword:
        query = query.filter(
            or_(
                User.name.contains(keyword),
                User.phone.contains(keyword),
                User.id_card.contains(keyword),
                Reservation.group_name.contains(keyword),
                Reservation.visiting_unit.contains(keyword)
            )
        )

    if status_filter:
        query = query.filter(Reservation.status == status_filter)

    if res_type_filter:
        query = query.filter(Reservation.res_type == res_type_filter)

    if start_date:
        query = query.filter(Reservation.visit_date >= start_date)

    if end_date:
        query = query.filter(Reservation.visit_date <= end_date)

    if visit_start:
        query = query.filter(func.date(Reservation.created_at) >= visit_start)

    if visit_end:
        query = query.filter(func.date(Reservation.created_at) <= visit_end)

    if venue_id:
        query = query.filter(Reservation.venue_id == int(venue_id))

    query = query.order_by(Reservation.created_at.desc())
    reservations = query.all()

    wb = Workbook()
    ws = wb.active
    ws.title = "预约记录"

    headers = [
        '序号', '申请时间', '预约类型', '申请人姓名', '证件类型', '证件号码',
        '手机号', '单位名称', '参观单位', '参观场馆', '校区', '参观日期',
        '参观时间', '参观人数', '需要讲解', '讲解员', '车牌号', '审批状态',
        '拒绝原因', '核销时间', '申请理由'
    ]

    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    for idx, r in enumerate(reservations, 1):
        row_data = [
            idx,
            r.created_at.strftime('%Y-%m-%d %H:%M') if r.created_at else '',
            '单位预约' if r.res_type == '单位' else '个人预约',
            r.user.name,
            r.user.id_type,
            r.user.id_card,
            r.user.phone,
            r.group_name or '',
            r.visiting_unit or '',
            r.venue.name,
            r.campus or '',
            r.visit_date.strftime('%Y-%m-%d') if r.visit_date else '',
            r.visit_time or '',
            r.group_size or r.visitor_count or 1,
            '是' if r.need_guide else '否',
            r.guide_info or '',
            r.license_plate or '',
            r.status,
            r.reject_reason or '',
            r.verified_at.strftime('%Y-%m-%d %H:%M') if r.verified_at else '',
            r.reason or '',
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=idx + 1, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 20
    ws.column_dimensions['G'].width = 14
    ws.column_dimensions['H'].width = 20
    ws.column_dimensions['I'].width = 20
    ws.column_dimensions['J'].width = 22
    ws.column_dimensions['K'].width = 14
    ws.column_dimensions['L'].width = 12
    ws.column_dimensions['M'].width = 14
    ws.column_dimensions['N'].width = 10
    ws.column_dimensions['O'].width = 10
    ws.column_dimensions['P'].width = 12
    ws.column_dimensions['Q'].width = 10
    ws.column_dimensions['R'].width = 10
    ws.column_dimensions['S'].width = 16
    ws.column_dimensions['T'].width = 18
    ws.column_dimensions['U'].width = 30

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'预约记录_{timestamp}.xlsx'
    )


@admin_bp.route("/api/stats")
def api_stats():
    period = request.args.get('period', 'today')
    custom_start = request.args.get('start', '').strip()
    custom_end = request.args.get('end', '').strip()
    categories = _get_area_categories()

    today = date.today()

    if custom_start and custom_end:
        start = datetime.strptime(custom_start, '%Y-%m-%d').date()
        end = datetime.strptime(custom_end, '%Y-%m-%d').date()
        period = 'custom'
    elif period == 'today':
        start = today
        end = today
    elif period == 'week':
        start = today - timedelta(days=today.weekday())
        end = today
    elif period == 'month':
        start = today.replace(day=1)
        end = today
    elif period == 'year':
        start = today.replace(month=1, day=1)
        end = today
    else:
        start = today
        end = today

    if start == end:
        date_cond = func.date(Reservation.created_at) == start
    else:
        date_cond = func.date(Reservation.created_at).between(start, end)

    venue_cond = Reservation.venue.has(Venue.category.in_(categories))

    total_count = db.session.query(func.count(Reservation.id)).filter(
        date_cond, venue_cond
    ).scalar() or 0

    pending_count = db.session.query(func.count(Reservation.id)).filter(
        date_cond, venue_cond, Reservation.status == '待审核'
    ).scalar() or 0

    approved_count = db.session.query(func.count(Reservation.id)).filter(
        date_cond, venue_cond, Reservation.status.in_(['已同意', '已核销'])
    ).scalar() or 0

    group_count = db.session.query(func.count(Reservation.id)).filter(
        date_cond, venue_cond, Reservation.res_type == '单位'
    ).scalar() or 0

    individual_count = db.session.query(func.count(Reservation.id)).filter(
        date_cond, venue_cond, Reservation.res_type == '个人'
    ).scalar() or 0

    verified_count = db.session.query(func.count(Reservation.id)).filter(
        date_cond, venue_cond, Reservation.status == '已核销'
    ).scalar() or 0

    rejected_count = db.session.query(func.count(Reservation.id)).filter(
        date_cond, venue_cond, Reservation.status == '已拒绝'
    ).scalar() or 0

    cancelled_count = db.session.query(func.count(Reservation.id)).filter(
        date_cond, venue_cond, Reservation.status == '已取消'
    ).scalar() or 0

    return jsonify({
        "period": period,
        "start_date": start.strftime('%Y-%m-%d'),
        "end_date": end.strftime('%Y-%m-%d'),
        "total_count": total_count,
        "pending_count": pending_count,
        "approved_count": approved_count,
        "group_count": group_count,
        "individual_count": individual_count,
        "verified_count": verified_count,
        "rejected_count": rejected_count,
        "cancelled_count": cancelled_count,
    })


@admin_bp.route("/config", methods=["POST"])
def config():
    config = SystemConfig.query.first()
    active_tab = request.form.get("active_tab", "config")

    if "toggle_system" in request.form:
        config.is_open = not config.is_open

    if "update_policy" in request.form:
        config.privacy_policy = request.form.get("privacy_policy")
        flash("隐私声明已更新")

    if "publish_notice" in request.form:
        title = request.form.get("title")
        content = request.form.get("content")
        if len(content) > 5000:
            flash("公告内容不能超过5000字符")
            return redirect(url_for("admin.dashboard", active_tab="notice"))
        new_notice = Announcement(title=title, content=content)
        db.session.add(new_notice)

    if "update_venue" in request.form:
        venue_id = request.form.get("venue_id", type=int)
        name = request.form.get("venue_name")
        description = request.form.get("venue_description")
        address = request.form.get("venue_address")
        campus = request.form.get("venue_campus")
        advance_days = request.form.get("venue_advance_days", 7, type=int)
        cutoff_time = request.form.get("venue_cutoff_time", "16:00")
        is_active = "venue_is_active" in request.form
        category = request.form.get("venue_category", "")
        open_hours = request.form.get("venue_open_hours", "09:00-11:00,14:00-16:00")

        venue = Venue.query.get(venue_id)
        if venue:
            venue.name = name
            venue.description = description
            venue.address = address
            venue.campus = campus
            venue.advance_days = advance_days
            venue.cutoff_time = cutoff_time
            venue.is_active = is_active
            if category:
                venue.category = category
            venue.open_hours = open_hours
            flash("场馆设置更新成功")
        else:
            flash("场馆不存在")

    if "add_venue" in request.form:
        name = request.form.get("venue_name")
        category = request.form.get("venue_category", "校史馆")
        campus = request.form.get("venue_campus", "")
        new_venue = Venue(
            name=name, category=category, campus=campus,
            address=request.form.get("venue_address", ""),
            open_hours=request.form.get("venue_open_hours", "09:00-11:00,14:00-16:00"),
            advance_days=request.form.get("venue_advance_days", 7, type=int),
            cutoff_time=request.form.get("venue_cutoff_time", "16:00"),
            description=request.form.get("venue_description", ""),
            is_active=True
        )
        db.session.add(new_venue)
        flash(f"场馆 {name} 已添加")

    if "delete_venue" in request.form:
        venue_id = request.form.get("venue_id", type=int)
        venue = Venue.query.get(venue_id)
        if venue:
            res_ids = [r.id for r in Reservation.query.with_entities(Reservation.id).filter_by(venue_id=venue_id).all()]
            CancelRequest.query.filter(CancelRequest.reservation_id.in_(res_ids)).delete(synchronize_session=False)
            Attachment.query.filter(Attachment.reservation_id.in_(res_ids)).delete(synchronize_session=False)
            Reservation.query.filter_by(venue_id=venue_id).delete()
            VenueTimeSlot.query.filter_by(venue_id=venue_id).delete()
            VenueTimeSlotDisabledDate.query.filter_by(venue_id=venue_id).delete()
            db.session.delete(venue)
            flash(f"场馆 {venue.name} 已删除")

    db.session.commit()
    return redirect(url_for("admin.dashboard", active_tab=active_tab))


@admin_bp.route("/handle-cancel-request/<int:req_id>", methods=["POST"])
def handle_cancel_request(req_id):
    action = request.form.get("action")
    remark = request.form.get("remark", "")

    cancel_req = db.session.get(CancelRequest, req_id)
    if not cancel_req:
        flash("撤销申请不存在")
        return redirect(url_for("admin.dashboard", active_tab="archive"))

    if cancel_req.status != '待处理':
        flash(f"该申请已被处理 (当前状态: {cancel_req.status})")
        return redirect(url_for("admin.dashboard", active_tab="archive"))

    reservation = db.session.get(Reservation, cancel_req.reservation_id)
    if not reservation:
        flash("关联的预约记录不存在")
        return redirect(url_for("admin.dashboard", active_tab="archive"))

    if action == 'approve':
        reservation_id = cancel_req.reservation_id

        cancel_req.status = "已同意"
        cancel_req.admin_remark = remark
        cancel_req.processed_at = datetime.now()

        other_cancel_requests = CancelRequest.query.filter(
            CancelRequest.reservation_id == reservation_id,
            CancelRequest.id != req_id
        ).all()
        for other_req in other_cancel_requests:
            other_req.status = "已拒绝"
            other_req.admin_remark = "因预约已被撤销，该申请被自动拒绝"
            other_req.processed_at = datetime.now()

        db.session.commit()

        db.session.execute(db.text("DELETE FROM attachment WHERE reservation_id = :rid"), {"rid": reservation_id})
        db.session.execute(db.text("DELETE FROM cancel_request WHERE reservation_id = :rid"), {"rid": reservation_id})
        db.session.execute(db.text("DELETE FROM reservation WHERE id = :rid"), {"rid": reservation_id})
        db.session.commit()

        logger.info("撤销申请 %s 已同意，预约 %s 已撤销，用户 %s",
                    req_id, reservation_id, cancel_req.user_id)
    elif action == 'reject':
        cancel_req.status = "已拒绝"
        cancel_req.admin_remark = remark
        cancel_req.processed_at = datetime.now()
        logger.info("撤销申请 %s 已拒绝，原因: %s", req_id, remark)
    else:
        flash("无效的操作")
        return redirect(url_for("admin.dashboard", active_tab="archive"))

    db.session.commit()
    flash(f"撤销申请已{'同意' if action == 'approve' else '拒绝'}")
    return redirect(url_for("admin.dashboard", active_tab="cancel"))


@admin_bp.route("/cancel-request/<int:req_id>")
def get_cancel_request_detail(req_id):
    if not session.get("admin_logged_in"):
        return jsonify({"error": "请先登录"}), 401

    cancel_req = db.session.get(CancelRequest, req_id)
    if not cancel_req:
        return jsonify({"error": "撤销申请不存在"}), 404

    reservation = db.session.get(Reservation, cancel_req.reservation_id)
    if not reservation:
        return jsonify({"error": "关联的预约记录不存在"}), 404

    venue = db.session.get(Venue, reservation.venue_id) if reservation.venue_id else None

    return jsonify({
        "cancel_request": {
            "id": cancel_req.id,
            "created_at": cancel_req.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            "status": cancel_req.status,
            "reason": cancel_req.reason,
            "admin_remark": cancel_req.admin_remark,
            "user_name": cancel_req.user.name,
            "user_id_type": cancel_req.user.id_type,
            "user_id_card": cancel_req.user.id_card,
            "user_phone": cancel_req.user.phone,
            "venue_name": venue.name if venue else "未知",
            "visit_date": reservation.visit_date.strftime('%Y-%m-%d') if reservation.visit_date else "",
            "visit_time": reservation.visit_time
        }
    })


@admin_bp.route("/announcement/<int:ann_id>/toggle-pin", methods=["POST"])
def toggle_announcement_pin(ann_id):
    announcement = db.session.get(Announcement, ann_id)
    if announcement:
        announcement.is_pinned = not announcement.is_pinned
        db.session.commit()
        flash(f"公告{'已顶置' if announcement.is_pinned else '已取消顶置'}")
    else:
        flash("公告不存在")
    return redirect(url_for("admin.dashboard", active_tab="notice"))


@admin_bp.route("/announcement/<int:ann_id>/toggle-hide", methods=["POST"])
def toggle_announcement_hide(ann_id):
    announcement = db.session.get(Announcement, ann_id)
    if announcement:
        announcement.is_hidden = not announcement.is_hidden
        db.session.commit()
        flash(f"公告{'已隐藏' if announcement.is_hidden else '已显示'}")
    else:
        flash("公告不存在")
    return redirect(url_for("admin.dashboard", active_tab="notice"))


@admin_bp.route("/announcement/<int:ann_id>/delete", methods=["POST"])
def delete_announcement(ann_id):
    announcement = db.session.get(Announcement, ann_id)
    if announcement:
        db.session.delete(announcement)
        db.session.commit()
        flash("公告已删除")
    else:
        flash("公告不存在")
    return redirect(url_for("admin.dashboard", active_tab="notice"))


@admin_bp.route("/announcement/<int:ann_id>/edit", methods=["POST"])
def edit_announcement(ann_id):
    announcement = db.session.get(Announcement, ann_id)
    if announcement:
        title = request.form.get("title")
        content = request.form.get("content")
        if len(content) > 5000:
            flash("公告内容不能超过5000字符")
            return redirect(url_for("admin.dashboard", active_tab="notice"))
        announcement.title = title
        announcement.content = content
        db.session.commit()
        flash("公告已更新")
    else:
        flash("公告不存在")
    return redirect(url_for("admin.dashboard", active_tab="notice"))


@admin_bp.route("/account", methods=["POST"])
def account():
    action = request.form.get("action")

    if action == "change_own_password":
        new_pass = request.form.get("new_password", "")
        if len(new_pass) < MIN_PASSWORD_LENGTH:
            flash(f"密码长度不能少于 {MIN_PASSWORD_LENGTH} 位")
            return redirect(url_for("admin.dashboard", active_tab="account"))
        admin_id = session.get("admin_id")
        current_admin = db.session.get(Admin, admin_id)
        if not current_admin:
            session.clear()
            flash("账号异常，请重新登录")
            return redirect(url_for("admin.login"))
        current_admin.password_hash = generate_password_hash(new_pass)
        db.session.commit()
        flash("您的密码已修改，请重新登录")
        session.clear()
        return redirect(url_for("admin.login"))

    if not session.get("is_super"):
        flash("无权操作")
        return redirect(url_for("admin.dashboard", active_tab="account"))

    if action == "create_admin":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        role_id = request.form.get("role_id", type=int)
        if len(password) < MIN_PASSWORD_LENGTH:
            flash(f"密码长度不能少于 {MIN_PASSWORD_LENGTH} 位")
            return redirect(url_for("admin.dashboard", active_tab="account"))
        if Admin.query.filter_by(username=username).first():
            flash("该用户名已存在")
        else:
            new_admin = Admin(
                username=username,
                password_hash=generate_password_hash(password),
                is_super=False,
                role_id=role_id if role_id else None
            )
            db.session.add(new_admin)
            db.session.commit()
            flash(f"管理员 {username} 创建成功")

    elif action == "edit_admin":
        admin_id = request.form.get("admin_id", type=int)
        role_id = request.form.get("role_id", type=int)
        target = db.session.get(Admin, admin_id)
        if target and not target.is_super:
            target.role_id = role_id if role_id else None
            db.session.commit()
            flash(f"管理员 {target.username} 角色已更新")
        else:
            flash("编辑失败")

    elif action == "delete_admin":
        admin_id = request.form.get("admin_id")
        target = db.session.get(Admin, admin_id)
        if target and not target.is_super:
            db.session.delete(target)
            db.session.commit()
            flash("管理员已删除")
        else:
            flash("删除失败：无法删除超级管理员或用户不存在")

    elif action == "reset_password":
        admin_id = request.form.get("admin_id")
        new_pass = request.form.get("new_password", "")
        if len(new_pass) < MIN_PASSWORD_LENGTH:
            flash(f"密码长度不能少于 {MIN_PASSWORD_LENGTH} 位")
            return redirect(url_for("admin.dashboard", active_tab="account"))
        target = db.session.get(Admin, admin_id)
        if target:
            target.password_hash = generate_password_hash(new_pass)
            db.session.commit()
            flash(f"管理员 {target.username} 的密码已重置")

    return redirect(url_for("admin.dashboard", active_tab="account"))


@admin_bp.route("/role", methods=["POST"])
def role_management():
    action = request.form.get("action")

    if action == "create":
        name = request.form.get("name", "").strip()
        if Role.query.filter_by(name=name).first():
            flash("该角色名称已存在")
        else:
            permissions = json.dumps(request.form.getlist("permissions"))
            role = Role(name=name, permissions=permissions)
            db.session.add(role)
            db.session.commit()
            flash(f"角色 {name} 创建成功")

    elif action == "edit":
        role_id = request.form.get("role_id", type=int)
        role = db.session.get(Role, role_id)
        if role:
            name = request.form.get("name", "").strip()
            existing = Role.query.filter(Role.name == name, Role.id != role_id).first()
            if existing:
                flash("该角色名称已存在")
            else:
                role.name = name
                role.permissions = json.dumps(request.form.getlist("permissions"))
                db.session.commit()
                flash(f"角色 {name} 已更新")
        else:
            flash("角色不存在")

    elif action == "delete":
        role_id = request.form.get("role_id", type=int)
        role = db.session.get(Role, role_id)
        if role:
            Admin.query.filter_by(role_id=role_id).update({Admin.role_id: None})
            db.session.delete(role)
            db.session.commit()
            flash(f"角色 {role.name} 已删除")

    return redirect(url_for("admin.dashboard", active_tab="roles"))


@admin_bp.route("/approval-staff", methods=["POST"])
def approval_staff_management():
    action = request.form.get("action")
    target_type = request.form.get("staff_type", "approval")

    if action == "create":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        name = request.form.get("name", "").strip()
        staff_id = request.form.get("staff_id", "").strip()
        department = request.form.get("department", "").strip()
        phone = request.form.get("phone", "").strip()
        venue_ids = json.dumps([int(v) for v in request.form.getlist("assigned_venues") if v])

        if not username:
            flash("用户名不能为空")
            return redirect(url_for("admin.dashboard", active_tab="approval_staff"))
        if len(password) < 8:
            flash("密码长度不能少于8位")
            return redirect(url_for("admin.dashboard", active_tab="approval_staff"))
        if ApprovalStaff.query.filter_by(username=username).first():
            flash("该用户名已存在")
            return redirect(url_for("admin.dashboard", active_tab="approval_staff"))

        staff = ApprovalStaff(
            username=username,
            password_hash=generate_password_hash(password),
            name=name, staff_id=staff_id, department=department,
            phone=phone, assigned_venue_ids=venue_ids,
            staff_type=target_type
        )
        db.session.add(staff)
        db.session.commit()
        type_label = '部门领导' if target_type == 'leader' else '审批教师'
        _write_log('operation', '人员管理', '添加人员', f'添加{type_label} {name}（{username}）')
        flash(f'{type_label} {name} 已添加，登录用户名：{username}')

    elif action == "edit":
        staff_id_pk = request.form.get("id", type=int)
        staff = db.session.get(ApprovalStaff, staff_id_pk)
        if staff:
            username = request.form.get("username", "").strip()
            if username and username != staff.username:
                if ApprovalStaff.query.filter(ApprovalStaff.username == username, ApprovalStaff.id != staff_id_pk).first():
                    flash("该用户名已存在")
                    return redirect(url_for("admin.dashboard", active_tab="approval_staff"))
                staff.username = username
            password = request.form.get("password", "")
            if password:
                if len(password) < 8:
                    flash("密码长度不能少于8位")
                    return redirect(url_for("admin.dashboard", active_tab="approval_staff"))
                staff.password_hash = generate_password_hash(password)
            staff.name = request.form.get("name", "").strip()
            staff.staff_id = request.form.get("staff_id", "").strip()
            staff.department = request.form.get("department", "").strip()
            staff.phone = request.form.get("phone", "").strip()
            staff.assigned_venue_ids = json.dumps([int(v) for v in request.form.getlist("assigned_venues") if v])
            staff.is_active = "is_active" in request.form
            if target_type:
                staff.staff_type = target_type
            db.session.commit()
            type_label = '部门领导' if staff.staff_type == 'leader' else '审批教师'
            _write_log('operation', '人员管理', '编辑人员', f'编辑{type_label} {staff.name}')
            flash(f'{type_label} {staff.name} 已更新')

    elif action == "delete":
        staff_id_pk = request.form.get("id", type=int)
        staff = db.session.get(ApprovalStaff, staff_id_pk)
        if staff:
            type_label = '部门领导' if staff.staff_type == 'leader' else '审批教师'
            _write_log('operation', '人员管理', '删除人员', f'删除{type_label} {staff.name}')
            db.session.delete(staff)
            db.session.commit()
            flash(f'{type_label} {staff.name} 已删除')

    elif action == "reset_password":
        staff_id_pk = request.form.get("id", type=int)
        new_pass = request.form.get("new_password", "")
        staff = db.session.get(ApprovalStaff, staff_id_pk)
        if staff:
            if len(new_pass) < 8:
                flash("密码长度不能少于8位")
                return redirect(url_for("admin.dashboard", active_tab="approval_staff"))
            staff.password_hash = generate_password_hash(new_pass)
            db.session.commit()
            type_label = '部门领导' if staff.staff_type == 'leader' else '审批教师'
            _write_log('operation', '人员管理', '重置密码', f'重置{type_label} {staff.name} 密码')
            flash(f'{type_label} {staff.name} 的密码已重置')

    return redirect(url_for("admin.dashboard", active_tab="approval_staff"))


@admin_bp.route("/guide", methods=["POST"])
def guide_management():
    action = request.form.get("action")

    if action == "create":
        name = request.form.get("name", "").strip()
        staff_id = request.form.get("staff_id", "").strip()
        phone = request.form.get("phone", "").strip()
        expertise = request.form.get("expertise", "").strip()
        status = request.form.get("status", "在岗")

        guide = Guide(
            name=name, staff_id=staff_id, phone=phone,
            expertise=expertise, status=status
        )
        db.session.add(guide)
        db.session.commit()
        flash(f"讲解员 {name} 已添加")

    elif action == "edit":
        guide_id = request.form.get("id", type=int)
        guide = db.session.get(Guide, guide_id)
        if guide:
            guide.name = request.form.get("name", "").strip()
            guide.staff_id = request.form.get("staff_id", "").strip()
            guide.phone = request.form.get("phone", "").strip()
            guide.expertise = request.form.get("expertise", "").strip()
            guide.status = request.form.get("status", "在岗")
            db.session.commit()
            flash(f"讲解员 {guide.name} 已更新")

    elif action == "toggle_status":
        guide_id = request.form.get("id", type=int)
        guide = db.session.get(Guide, guide_id)
        if guide:
            guide.status = '离岗' if guide.status == '在岗' else '在岗'
            db.session.commit()
            flash(f"讲解员 {guide.name} 状态已切换为 {guide.status}")

    elif action == "delete":
        guide_id = request.form.get("id", type=int)
        guide = db.session.get(Guide, guide_id)
        if guide:
            db.session.delete(guide)
            db.session.commit()
            flash(f"讲解员 {guide.name} 已删除")

    return redirect(url_for("admin.dashboard", active_tab="guides"))


@admin_bp.route("/user", methods=["POST"])
def user_management():
    action = request.form.get("action")

    if action == "toggle_blacklist":
        user_id = request.form.get("user_id", type=int)
        user = db.session.get(User, user_id)
        if user:
            user.is_blacklisted = not user.is_blacklisted
            db.session.commit()
            flash(f"用户 {user.name} {'已加入黑名单' if user.is_blacklisted else '已从黑名单移除'}")

    return redirect(url_for("admin.dashboard", active_tab="users"))


@admin_bp.route("/api/users")
def api_users():
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 15, type=int)
    keyword = request.args.get('keyword', '').strip()
    blacklist_filter = request.args.get('blacklist', '').strip()

    query = User.query

    if keyword:
        query = query.filter(
            or_(
                User.name.contains(keyword),
                User.id_card.contains(keyword),
                User.phone.contains(keyword)
            )
        )

    if blacklist_filter == '1':
        query = query.filter_by(is_blacklisted=True)
    elif blacklist_filter == '0':
        query = query.filter_by(is_blacklisted=False)

    query = query.order_by(User.created_at.desc())

    pagination = query.paginate(page=page, per_page=per_page, error_out=False)

    return jsonify({
        "users": [
            {
                "id": u.id,
                "name": u.name,
                "id_type": u.id_type,
                "id_card": u.id_card,
                "phone": u.phone,
                "email": u.email or '',
                "is_blacklisted": u.is_blacklisted,
                "reservation_count": len(u.reservations),
                "created_at": u.created_at.strftime('%Y-%m-%d %H:%M') if u.created_at else '',
            }
            for u in pagination.items
        ],
        "pagination": {
            "page": pagination.page,
            "pages": pagination.pages,
            "has_prev": pagination.has_prev,
            "has_next": pagination.has_next,
            "total": pagination.total
        }
    })


@admin_bp.route("/api/user/<int:user_id>/reservations")
def api_user_reservations(user_id):
    reservations = Reservation.query.filter_by(user_id=user_id).join(Venue).order_by(
        Reservation.created_at.desc()
    ).all()

    return jsonify({
        "reservations": [
            {
                "id": r.id,
                "venue_name": r.venue.name,
                "venue_category": r.venue.category,
                "visit_date": r.visit_date.strftime('%Y-%m-%d') if r.visit_date else '',
                "visit_time": r.visit_time or '',
                "res_type": r.res_type,
                "status": r.status,
                "group_name": r.group_name or '',
                "visitor_count": r.visitor_count or 1,
                "created_at": r.created_at.strftime('%Y-%m-%d %H:%M') if r.created_at else '',
            }
            for r in reservations
        ]
    })


@admin_bp.route("/logout")
def logout():
    _write_log('login', '系统退出', '管理员退出', f'{session.get("admin_username", "")} 安全退出')
    session.clear()
    flash("您已安全退出")
    return redirect(url_for("admin.login"))


@admin_bp.route("/stats")
def stats():
    categories = _get_area_categories()
    venues = Venue.query.filter(
        Venue.category.in_(categories)
    ).all()
    return render_template(
        "admin_stats.html",
        venues=venues,
        management_area=_get_management_area(),
        area_label=_get_area_label(),
        categories=json.dumps(categories),
    )


@admin_bp.route("/time-slot-config", methods=["POST"])
def time_slot_config():
    venue_id = request.form.get("venue_id", type=int)
    active_tab = request.form.get("active_tab", "venue")
    if not venue_id:
        flash("场馆ID无效")
        return redirect(url_for("admin.dashboard", active_tab=active_tab))

    form_data = request.form

    time_slot_values = set()
    for key in form_data:
        if key.startswith('individual_capacity_'):
            parts = key.split('_')
            if len(parts) > 3:
                time_slot = '_'.join(parts[2:-1])
                time_slot_values.add(time_slot)

    VenueTimeSlot.query.filter_by(venue_id=venue_id).delete()

    for time_slot in time_slot_values:
        for day in range(7):
            individual_capacity_key = f"individual_capacity_{time_slot}_{day}"
            active_key = f"active_{time_slot}_{day}"
            is_group_active_key = f"is_group_active_{time_slot}_{day}"

            individual_capacity = request.form.get(individual_capacity_key, type=int)
            is_active = active_key in request.form
            is_group_active = is_group_active_key in request.form

            if individual_capacity:
                new_slot = VenueTimeSlot(
                    venue_id=venue_id,
                    day_of_week=day,
                    time_slot=time_slot,
                    individual_capacity=individual_capacity,
                    is_group_active=is_group_active,
                    is_active=is_active
                )
                db.session.add(new_slot)

    db.session.commit()
    flash("时段设置更新成功")
    return redirect(url_for("admin.dashboard", active_tab=active_tab))


@admin_bp.route("/get-time-slots/<int:venue_id>")
def get_time_slots(venue_id):
    time_slots = VenueTimeSlot.query.filter_by(venue_id=venue_id).all()

    slot_data = {}
    for slot in time_slots:
        if slot.time_slot not in slot_data:
            slot_data[slot.time_slot] = {}
        slot_data[slot.time_slot][slot.day_of_week] = {
            "individual_capacity": slot.individual_capacity,
            "is_active": slot.is_active,
            "is_group_active": slot.is_group_active
        }

    return {"slots": slot_data}


@admin_bp.route("/get-attachments/<int:res_id>")
def get_attachments(res_id):
    reservation = db.session.get(Reservation, res_id)
    if not reservation:
        return {"attachments": []}

    attachments = []
    for attachment in reservation.attachments:
        attachments.append({
            "filename": attachment.filename,
            "filepath": attachment.filepath,
            "file_size": attachment.file_size
        })

    return {"attachments": attachments}


@admin_bp.route("/get-venues")
def get_venues():
    categories = _get_area_categories()
    venues = Venue.query.filter(
        Venue.category.in_(categories)
    ).all()

    venue_list = []
    for venue in venues:
        venue_list.append({
            "id": venue.id,
            "name": venue.name,
            "category": venue.category,
            "campus": venue.campus or '',
        })

    return {"venues": venue_list}


@admin_bp.route("/get-all-venues")
def get_all_venues():
    categories = _get_area_categories()
    venues = Venue.query.filter(
        Venue.category.in_(categories)
    ).all()

    return jsonify({
        "venues": [
            {"id": v.id, "name": v.name, "category": v.category, "campus": v.campus or ''}
            for v in venues
        ]
    })


@admin_bp.route("/stats/reservation-trend")
def reservation_trend():
    time_range = request.args.get("time_range", "month")
    categories = _get_area_categories()

    query = Reservation.query.join(Venue).filter(Venue.category.in_(categories))

    if time_range == "day":
        result = db.session.query(
            func.date(Reservation.created_at).label('date'),
            func.count(Reservation.id).label('count')
        ).filter(Reservation.venue.has(Venue.category.in_(categories))).group_by(
            func.date(Reservation.created_at)
        ).order_by('date').all()
        data = {
            "labels": [item.date.strftime('%Y-%m-%d') for item in result],
            "values": [item.count for item in result]
        }
    elif time_range == "week":
        result = db.session.query(
            func.date_format(Reservation.created_at, '%Y-%u').label('week'),
            func.count(Reservation.id).label('count')
        ).filter(Reservation.venue.has(Venue.category.in_(categories))).group_by(
            'week'
        ).order_by('week').all()
        data = {
            "labels": [item.week for item in result],
            "values": [item.count for item in result]
        }
    else:
        result = db.session.query(
            func.date_format(Reservation.created_at, '%Y-%m').label('month'),
            func.count(Reservation.id).label('count')
        ).filter(Reservation.venue.has(Venue.category.in_(categories))).group_by(
            'month'
        ).order_by('month').all()
        data = {
            "labels": [item.month for item in result],
            "values": [item.count for item in result]
        }
    return data


@admin_bp.route("/stats/reservation-status")
def reservation_status():
    categories = _get_area_categories()
    result = db.session.query(
        Reservation.status,
        func.count(Reservation.id).label('count')
    ).join(Venue).filter(Venue.category.in_(categories)).group_by(
        Reservation.status
    ).all()
    data = {
        "labels": [item.status for item in result],
        "values": [item.count for item in result]
    }
    return data


@admin_bp.route("/stats/venue-comparison")
def venue_comparison():
    categories = _get_area_categories()
    result = db.session.query(
        Venue.name,
        func.count(Reservation.id).label('count')
    ).join(Reservation, Venue.id == Reservation.venue_id).filter(
        Venue.category.in_(categories)
    ).group_by(Venue.name).all()
    data = {
        "labels": [item.name for item in result],
        "values": [item.count for item in result]
    }
    return data


@admin_bp.route("/stats/reservation-type")
def reservation_type():
    categories = _get_area_categories()
    result = db.session.query(
        Reservation.res_type,
        func.count(Reservation.id).label('count')
    ).join(Venue).filter(Venue.category.in_(categories)).group_by(
        Reservation.res_type
    ).all()
    data = {
        "labels": [item.res_type for item in result],
        "values": [item.count for item in result]
    }
    return data


@admin_bp.route("/stats/visitor-trend")
def visitor_trend():
    time_range = request.args.get("time_range", "month")
    categories = _get_area_categories()

    if time_range == "day":
        result = db.session.query(
            func.date(Reservation.visit_date).label('date'),
            func.sum(Reservation.group_size).label('count')
        ).join(Venue).filter(
            Reservation.status == "已同意",
            Venue.category.in_(categories)
        ).group_by(func.date(Reservation.visit_date)).order_by('date').all()
        data = {
            "labels": [item.date.strftime('%Y-%m-%d') for item in result],
            "values": [int(item.count or 0) for item in result]
        }
    elif time_range == "week":
        result = db.session.query(
            func.date_format(Reservation.visit_date, '%Y-%u').label('week'),
            func.sum(Reservation.group_size).label('count')
        ).join(Venue).filter(
            Reservation.status == "已同意",
            Venue.category.in_(categories)
        ).group_by('week').order_by('week').all()
        data = {
            "labels": [item.week for item in result],
            "values": [int(item.count or 0) for item in result]
        }
    else:
        result = db.session.query(
            func.date_format(Reservation.visit_date, '%Y-%m').label('month'),
            func.sum(Reservation.group_size).label('count')
        ).join(Venue).filter(
            Reservation.status == "已同意",
            Venue.category.in_(categories)
        ).group_by('month').order_by('month').all()
        data = {
            "labels": [item.month for item in result],
            "values": [int(item.count or 0) for item in result]
        }
    return data


@admin_bp.route("/stats/peak-hours")
def peak_hours():
    categories = _get_area_categories()
    result = db.session.query(
        Reservation.visit_time,
        func.sum(Reservation.group_size).label('count')
    ).join(Venue).filter(
        Reservation.status == "已同意",
        Venue.category.in_(categories)
    ).group_by(Reservation.visit_time).order_by('count').all()
    data = {
        "labels": [item.visit_time for item in result],
        "values": [int(item.count or 0) for item in result]
    }
    return data


@admin_bp.route("/stats/campus-comparison")
def campus_comparison():
    categories = _get_area_categories()
    result = db.session.query(
        Venue.campus,
        func.sum(Reservation.group_size).label('count')
    ).join(Reservation, Venue.id == Reservation.venue_id).filter(
        Reservation.status == "已同意",
        Venue.category.in_(categories)
    ).group_by(Venue.campus).all()
    data = {
        "labels": [item.campus for item in result],
        "values": [int(item.count or 0) for item in result]
    }
    return data


@admin_bp.route("/stats/time-utilization")
def time_utilization():
    categories = _get_area_categories()
    venue_ids = [v.id for v in Venue.query.filter(
        Venue.category.in_(categories)
    ).all()]

    time_slots = VenueTimeSlot.query.filter(
        VenueTimeSlot.venue_id.in_(venue_ids)
    ).distinct(VenueTimeSlot.time_slot).all()
    time_slot_list = [slot.time_slot for slot in time_slots]

    utilization_data = []
    for time_slot in time_slot_list:
        capacity = db.session.query(
            func.sum(VenueTimeSlot.individual_capacity)
        ).filter(
            VenueTimeSlot.time_slot == time_slot,
            VenueTimeSlot.venue_id.in_(venue_ids)
        ).scalar() or 0

        usage = db.session.query(
            func.sum(Reservation.group_size)
        ).join(Venue).filter(
            Reservation.visit_time == time_slot,
            Reservation.status == "已同意",
            Venue.category.in_(categories)
        ).scalar() or 0

        utilization = (usage / capacity * 100) if capacity > 0 else 0

        utilization_data.append({
            "time_slot": time_slot,
            "utilization": round(utilization, 2)
        })

    utilization_data.sort(key=lambda x: x["utilization"], reverse=True)

    data = {
        "labels": [item["time_slot"] for item in utilization_data],
        "values": [item["utilization"] for item in utilization_data]
    }

    return data


@admin_bp.route("/stats/total")
def total_statistics():
    categories = _get_area_categories()

    total_reservations = db.session.query(
        func.count(Reservation.id)
    ).join(Venue).filter(Venue.category.in_(categories)).scalar() or 0

    approved_reservations = db.session.query(
        func.count(Reservation.id)
    ).join(Venue).filter(
        Reservation.status == "已同意",
        Venue.category.in_(categories)
    ).scalar() or 0

    total_visitors = db.session.query(
        func.sum(Reservation.group_size)
    ).join(Venue).filter(
        Reservation.status == "已同意",
        Venue.category.in_(categories)
    ).scalar() or 0

    venue_ids = [v.id for v in Venue.query.filter(
        Venue.category.in_(categories)
    ).all()]

    total_capacity = db.session.query(
        func.sum(VenueTimeSlot.individual_capacity)
    ).filter(VenueTimeSlot.venue_id.in_(venue_ids)).scalar() or 0

    total_usage = db.session.query(
        func.sum(Reservation.group_size)
    ).join(Venue).filter(
        Reservation.status == "已同意",
        Venue.category.in_(categories)
    ).scalar() or 0

    avg_utilization = (total_usage / total_capacity * 100) if total_capacity > 0 else 0

    data = {
        "total_reservations": total_reservations,
        "approved_reservations": approved_reservations,
        "total_visitors": int(total_visitors) if total_visitors else 0,
        "avg_utilization": round(avg_utilization, 2)
    }

    return data


@admin_bp.route("/archive-requests")
def archive_requests():
    status_filter = request.args.get('status', '').strip()
    page = request.args.get('page', 1, type=int)

    query = ArchiveRequest.query.join(User)

    if status_filter:
        query = query.filter(ArchiveRequest.status == status_filter)

    query = query.order_by(ArchiveRequest.created_at.desc())

    pagination = query.paginate(page=page, per_page=10, error_out=False)
    archive_requests_list = pagination.items

    return jsonify({
        "archive_requests": [
            {
                "id": req.id,
                "user_id": req.user_id,
                "user_name": req.user.name,
                "user_id_type": req.user.id_type,
                "user_id_card": req.user.id_card,
                "user_phone": req.user.phone,
                "user_email": req.user.email,
                "request_email": req.email,
                "status": req.status,
                "admin_remark": req.admin_remark,
                "created_at": req.created_at.strftime('%Y-%m-%d %H:%M') if req.created_at else '',
                "updated_at": req.updated_at.strftime('%Y-%m-%d %H:%M') if req.updated_at else ''
            }
            for req in archive_requests_list
        ],
        "pagination": {
            "page": pagination.page,
            "pages": pagination.pages,
            "has_prev": pagination.has_prev,
            "has_next": pagination.has_next,
            "total": pagination.total
        },
        "curr_status": status_filter
    })


@admin_bp.route("/archive-request/<int:req_id>", methods=["POST"])
def handle_archive_request(req_id):
    action = request.form.get("action")
    remark = request.form.get("remark", "")

    archive_req = db.session.get(ArchiveRequest, req_id)
    if not archive_req:
        return jsonify({"error": "申请记录不存在"}), 404

    if archive_req.status != '待处理':
        return jsonify({"error": f"该申请已被处理 (当前状态: {archive_req.status})"}), 400

    if action == 'approve':
        archive_req.status = "已处理"
        archive_req.admin_remark = remark
        logger.info("档案查询申请 %s 已处理，用户 %s", req_id, archive_req.user_id)
    elif action == 'reject':
        archive_req.status = "已拒绝"
        archive_req.admin_remark = remark
        logger.info("档案查询申请 %s 已拒绝，原因: %s", req_id, remark)
    else:
        return jsonify({"error": "无效的操作"}), 400

    db.session.commit()
    flash(f"申请已{'处理' if action == 'approve' else '拒绝'}")
    return redirect(url_for("admin.dashboard", active_tab="archive"))


@admin_bp.route("/get-disabled-dates/<int:venue_id>")
def get_disabled_dates(venue_id):
    disabled_dates = VenueTimeSlotDisabledDate.query.filter_by(venue_id=venue_id).all()
    disabled_list = []
    for item in disabled_dates:
        disabled_list.append({
            "id": item.id,
            "time_slot": item.time_slot,
            "disabled_date": item.disabled_date.strftime('%Y-%m-%d')
        })
    return {"disabled_dates": disabled_list}


@admin_bp.route("/disabled-date", methods=["POST"])
def manage_disabled_date():
    action = request.form.get("action")
    venue_id = request.form.get("venue_id", type=int)
    time_slot = request.form.get("time_slot", "")
    disabled_date = request.form.get("disabled_date", "")
    active_tab = request.form.get("active_tab", "venue")

    if not venue_id:
        flash("场馆ID无效")
        return redirect(url_for("admin.dashboard", active_tab=active_tab))

    if action == "add":
        if not time_slot or not disabled_date:
            flash("请选择时段和日期")
            return redirect(url_for("admin.dashboard", active_tab=active_tab))

        existing = VenueTimeSlotDisabledDate.query.filter_by(
            venue_id=venue_id,
            time_slot=time_slot,
            disabled_date=disabled_date
        ).first()

        if existing:
            flash("该日期时段已被禁用")
            return redirect(url_for("admin.dashboard", active_tab=active_tab))

        new_disabled = VenueTimeSlotDisabledDate(
            venue_id=venue_id,
            time_slot=time_slot,
            disabled_date=disabled_date
        )
        db.session.add(new_disabled)
        flash("禁用日期添加成功")

    elif action == "delete":
        disabled_id = request.form.get("disabled_id", type=int)
        if disabled_id:
            disabled_record = db.session.get(VenueTimeSlotDisabledDate, disabled_id)
            if disabled_record:
                db.session.delete(disabled_record)
                flash("禁用日期已删除")
        else:
            flash("禁用记录ID无效")

    db.session.commit()
    return redirect(url_for("admin.dashboard", active_tab=active_tab))


@admin_bp.route("/switch-area/<area>")
def switch_area(area):
    if area in ('xiaoshi', 'archive'):
        session['management_area'] = area
    return redirect(url_for('admin.dashboard'))


@admin_bp.route("/api/reservation/<int:res_id>")
def api_reservation_detail(res_id):
    res = db.session.get(Reservation, res_id)
    if not res:
        return jsonify({"error": "预约记录不存在"}), 404

    attachments = []
    for att in res.attachments:
        attachments.append({
            "filename": att.filename,
            "filepath": att.filepath,
            "file_size": att.file_size
        })

    cancel_requests = CancelRequest.query.filter_by(reservation_id=res_id).order_by(
        CancelRequest.created_at.desc()
    ).all()
    cancel_list = [{
        "id": cr.id,
        "reason": cr.reason or '',
        "status": cr.status,
        "admin_remark": cr.admin_remark or '',
        "created_at": cr.created_at.strftime('%Y-%m-%d %H:%M') if cr.created_at else '',
        "processed_at": cr.processed_at.strftime('%Y-%m-%d %H:%M') if cr.processed_at else ''
    } for cr in cancel_requests]

    return jsonify({
        "id": res.id,
        "user_name": res.user.name,
        "user_phone": res.user.phone,
        "user_email": res.user.email or '',
        "user_id_type": res.user.id_type,
        "user_id_card": res.user.id_card,
        "identity": res.identity or '',
        "res_type": res.res_type,
        "group_name": res.group_name or '',
        "group_contact": res.group_contact or '',
        "group_size": res.group_size or 0,
        "visiting_unit": res.visiting_unit or '',
        "contact_phone": res.contact_phone or '',
        "visitor_count": res.visitor_count or 1,
        "license_plate": res.license_plate or '',
        "need_guide": res.need_guide,
        "visit_date": res.visit_date.strftime('%Y-%m-%d') if res.visit_date else '',
        "visit_time": res.visit_time or '',
        "venue_name": res.venue.name,
        "venue_category": res.venue.category,
        "campus": res.campus or '',
        "reason": res.reason or '',
        "archive_name": res.archive_name or '',
        "archive_number": res.archive_number or '',
        "archive_purpose": res.archive_purpose or '',
        "visit_type": res.visit_type or '线下',
        "status": res.status,
        "reject_reason": res.reject_reason or '',
        "guide_info": res.guide_info or '',
        "verified_at": res.verified_at.strftime('%Y-%m-%d %H:%M') if res.verified_at else '',
        "created_at": res.created_at.strftime('%Y-%m-%d %H:%M'),
        "attachments": attachments,
        "cancel_requests": cancel_list,
    })


@admin_bp.route("/api/reservation/<int:res_id>/delete", methods=["POST"])
def api_delete_reservation(res_id):
    res = db.session.get(Reservation, res_id)
    if not res:
        return jsonify({"error": "预约记录不存在"}), 404

    venue_name = res.venue.name if res.venue else '未知'
    user_name = res.user.name if res.user else '未知'

    db.session.execute(db.text("DELETE FROM attachment WHERE reservation_id = :rid"), {"rid": res_id})
    db.session.execute(db.text("DELETE FROM cancel_request WHERE reservation_id = :rid"), {"rid": res_id})
    db.session.execute(db.text("DELETE FROM reservation WHERE id = :rid"), {"rid": res_id})
    db.session.commit()

    _write_log('operation', '预约记录管理', '删除预约',
               f'删除预约记录 #{res_id}：{user_name} - {venue_name}')
    logger.info("管理员删除预约记录 %s，用户 %s，场馆 %s", res_id, user_name, venue_name)
    flash("预约记录已删除")
    return jsonify({"success": True})


@admin_bp.route("/api/users/export")
def export_users():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    keyword = request.args.get('keyword', '').strip()
    blacklist_filter = request.args.get('blacklist', '').strip()

    query = User.query
    if keyword:
        query = query.filter(
            or_(User.name.contains(keyword), User.id_card.contains(keyword), User.phone.contains(keyword))
        )
    if blacklist_filter == '1':
        query = query.filter_by(is_blacklisted=True)
    elif blacklist_filter == '0':
        query = query.filter_by(is_blacklisted=False)

    query = query.order_by(User.created_at.desc())
    users = query.all()

    wb = Workbook()
    ws = wb.active
    ws.title = "申请人信息"

    headers = ['序号', '姓名', '证件类型', '证件号码', '手机号', '邮箱', '黑名单', '预约次数', '注册时间']
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    for idx, u in enumerate(users, 1):
        row_data = [
            idx, u.name, u.id_type, u.id_card, u.phone,
            u.email or '', '是' if u.is_blacklisted else '否',
            len(u.reservations),
            u.created_at.strftime('%Y-%m-%d %H:%M') if u.created_at else ''
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=idx + 1, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')

    for col_letter, width in [('A', 6), ('B', 10), ('C', 10), ('D', 22), ('E', 14), ('F', 26), ('G', 8), ('H', 10), ('I', 18)]:
        ws.column_dimensions[col_letter].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    _write_log('operation', '用户管理', '导出用户', f'导出用户列表，共 {len(users)} 条')
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'申请人信息_{timestamp}.xlsx'
    )


@admin_bp.route("/api/users/delete", methods=["POST"])
def api_delete_user():
    data = request.get_json()
    user_id = data.get("user_id") if data else request.form.get("user_id", type=int)
    user = db.session.get(User, user_id)
    if not user:
        return jsonify({"error": "用户不存在"}), 404

    user_name = user.name
    reservations = Reservation.query.filter_by(user_id=user_id).all()
    for res in reservations:
        db.session.execute(db.text("DELETE FROM attachment WHERE reservation_id = :rid"), {"rid": res.id})
        db.session.execute(db.text("DELETE FROM cancel_request WHERE reservation_id = :rid"), {"rid": res.id})
    ArchiveRequest.query.filter_by(user_id=user_id).delete()
    Reservation.query.filter_by(user_id=user_id).delete()
    db.session.delete(user)
    db.session.commit()

    _write_log('operation', '用户管理', '删除用户', f'删除用户 {user_name}（#{user_id}）')
    flash(f"用户 {user_name} 已删除")
    return jsonify({"success": True})


@admin_bp.route("/stats/custom-report")
def custom_report():
    from sqlalchemy import extract
    categories = _get_area_categories()
    group_by = request.args.get('group_by', 'month')
    start_date = request.args.get('start', '').strip()
    end_date = request.args.get('end', '').strip()
    venue_id = request.args.get('venue_id', '').strip()
    status_filter = request.args.get('status', '').strip()
    res_type_filter = request.args.get('res_type', '').strip()

    query = db.session.query(Reservation).join(Venue).filter(
        Venue.category.in_(categories)
    )

    if start_date:
        query = query.filter(func.date(Reservation.created_at) >= start_date)
    if end_date:
        query = query.filter(func.date(Reservation.created_at) <= end_date)
    if venue_id:
        query = query.filter(Reservation.venue_id == int(venue_id))
    if status_filter:
        query = query.filter(Reservation.status == status_filter)
    if res_type_filter:
        query = query.filter(Reservation.res_type == res_type_filter)

    if group_by == 'day':
        result = db.session.query(
            func.date(Reservation.created_at).label('label'),
            func.count(Reservation.id).label('count')
        ).select_from(Reservation).join(Venue).filter(
            Venue.category.in_(categories)
        )
        if start_date:
            result = result.filter(func.date(Reservation.created_at) >= start_date)
        if end_date:
            result = result.filter(func.date(Reservation.created_at) <= end_date)
        if venue_id:
            result = result.filter(Reservation.venue_id == int(venue_id))
        if status_filter:
            result = result.filter(Reservation.status == status_filter)
        if res_type_filter:
            result = result.filter(Reservation.res_type == res_type_filter)
        result = result.group_by(func.date(Reservation.created_at)).order_by('label').all()
    elif group_by == 'week':
        result = db.session.query(
            func.date_format(Reservation.created_at, '%Y-%u').label('label'),
            func.count(Reservation.id).label('count')
        ).select_from(Reservation).join(Venue).filter(
            Venue.category.in_(categories)
        )
        if start_date:
            result = result.filter(func.date(Reservation.created_at) >= start_date)
        if end_date:
            result = result.filter(func.date(Reservation.created_at) <= end_date)
        if venue_id:
            result = result.filter(Reservation.venue_id == int(venue_id))
        if status_filter:
            result = result.filter(Reservation.status == status_filter)
        if res_type_filter:
            result = result.filter(Reservation.res_type == res_type_filter)
        result = result.group_by('label').order_by('label').all()
    else:
        result = db.session.query(
            func.date_format(Reservation.created_at, '%Y-%m').label('label'),
            func.count(Reservation.id).label('count')
        ).select_from(Reservation).join(Venue).filter(
            Venue.category.in_(categories)
        )
        if start_date:
            result = result.filter(func.date(Reservation.created_at) >= start_date)
        if end_date:
            result = result.filter(func.date(Reservation.created_at) <= end_date)
        if venue_id:
            result = result.filter(Reservation.venue_id == int(venue_id))
        if status_filter:
            result = result.filter(Reservation.status == status_filter)
        if res_type_filter:
            result = result.filter(Reservation.res_type == res_type_filter)
        result = result.group_by('label').order_by('label').all()

    total = query.count()
    return jsonify({
        "labels": [r.label for r in result],
        "values": [r.count for r in result],
        "total": total,
    })


@admin_bp.route("/stats/export-report")
def export_stats_report():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    categories = _get_area_categories()
    group_by = request.args.get('group_by', 'month')
    start_date = request.args.get('start', '').strip()
    end_date = request.args.get('end', '').strip()
    venue_id = request.args.get('venue_id', '').strip()
    status_filter = request.args.get('status', '').strip()
    res_type_filter = request.args.get('res_type', '').strip()

    result = db.session.query(
        func.date_format(Reservation.created_at, '%Y-%m').label('label'),
        func.count(Reservation.id).label('count')
    ).select_from(Reservation).join(Venue).filter(
        Venue.category.in_(categories)
    )
    if start_date:
        result = result.filter(func.date(Reservation.created_at) >= start_date)
    if end_date:
        result = result.filter(func.date(Reservation.created_at) <= end_date)
    if venue_id:
        result = result.filter(Reservation.venue_id == int(venue_id))
    if status_filter:
        result = result.filter(Reservation.status == status_filter)
    if res_type_filter:
        result = result.filter(Reservation.res_type == res_type_filter)
    result = result.group_by('label').order_by('label').all()

    wb = Workbook()
    ws = wb.active
    ws.title = "统计报表"

    headers = ['时间维度', '预约数量']
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    for idx, r in enumerate(result, 1):
        for col, val in enumerate([r.label, r.count], 1):
            cell = ws.cell(row=idx + 1, column=col, value=val)
            cell.border = thin_border

    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 14

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    _write_log('operation', '数据统计', '导出报表', f'导出统计报表')
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'统计报表_{timestamp}.xlsx'
    )


@admin_bp.route("/api/logs")
def api_logs():
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 20, type=int)
    log_type = request.args.get('log_type', '').strip()
    keyword = request.args.get('keyword', '').strip()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()

    query = SystemLog.query

    if log_type:
        query = query.filter(SystemLog.log_type == log_type)

    if keyword:
        query = query.filter(
            or_(
                SystemLog.operator.contains(keyword),
                SystemLog.action.contains(keyword),
                SystemLog.detail.contains(keyword),
                SystemLog.module.contains(keyword)
            )
        )

    if start_date:
        query = query.filter(func.date(SystemLog.created_at) >= start_date)
    if end_date:
        query = query.filter(func.date(SystemLog.created_at) <= end_date)

    query = query.order_by(SystemLog.created_at.desc())
    pagination = query.paginate(page=page, per_page=per_page, error_out=False)

    log_type_map = {'operation': '操作日志', 'login': '登录日志', 'error': '错误日志', 'access': '通行权限下发日志'}
    return jsonify({
        "logs": [
            {
                "id": log.id,
                "log_type": log.log_type,
                "log_type_label": log_type_map.get(log.log_type, log.log_type),
                "operator": log.operator,
                "operator_ip": log.operator_ip,
                "module": log.module,
                "action": log.action,
                "detail": log.detail or '',
                "result": log.result,
                "created_at": log.created_at.strftime('%Y-%m-%d %H:%M:%S') if log.created_at else ''
            }
            for log in pagination.items
        ],
        "pagination": {
            "page": pagination.page,
            "pages": pagination.pages,
            "has_prev": pagination.has_prev,
            "has_next": pagination.has_next,
            "total": pagination.total
        }
    })


@admin_bp.route("/api/logs/export")
def export_logs():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    log_type = request.args.get('log_type', '').strip()
    keyword = request.args.get('keyword', '').strip()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()

    query = SystemLog.query

    if log_type:
        query = query.filter(SystemLog.log_type == log_type)
    if keyword:
        query = query.filter(
            or_(
                SystemLog.operator.contains(keyword),
                SystemLog.action.contains(keyword),
                SystemLog.detail.contains(keyword),
                SystemLog.module.contains(keyword)
            )
        )
    if start_date:
        query = query.filter(func.date(SystemLog.created_at) >= start_date)
    if end_date:
        query = query.filter(func.date(SystemLog.created_at) <= end_date)

    query = query.order_by(SystemLog.created_at.desc())
    logs = query.all()

    log_type_map = {'operation': '操作日志', 'login': '登录日志', 'error': '错误日志', 'access': '通行权限下发日志'}

    wb = Workbook()
    ws = wb.active
    ws.title = "系统日志"

    headers = ['序号', '日志类型', '操作时间', '操作人', 'IP地址', '操作模块', '操作动作', '详细内容', '结果']
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    for idx, log in enumerate(logs, 1):
        row_data = [
            idx,
            log_type_map.get(log.log_type, log.log_type),
            log.created_at.strftime('%Y-%m-%d %H:%M:%S') if log.created_at else '',
            log.operator, log.operator_ip or '', log.module,
            log.action, log.detail or '', log.result
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=idx + 1, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')

    for col_letter, width in [('A', 6), ('B', 14), ('C', 20), ('D', 12), ('E', 16), ('F', 12), ('G', 14), ('H', 40), ('I', 8)]:
        ws.column_dimensions[col_letter].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    _write_log('operation', '日志管理', '导出日志', f'导出系统日志')
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'系统日志_{timestamp}.xlsx'
    )


@admin_bp.route("/api/logs/clear", methods=["POST"])
def clear_logs():
    days = request.form.get('days', 90, type=int)
    cutoff_date = datetime.now() - timedelta(days=days)
    count = SystemLog.query.filter(SystemLog.created_at < cutoff_date).delete()
    db.session.commit()
    _write_log('operation', '日志管理', '清理日志', f'清理 {days} 天前的日志，共 {count} 条')
    flash(f"已清理 {count} 条 {days} 天前的日志")
    return jsonify({"success": True, "count": count})


@admin_bp.route("/config/system", methods=["POST"])
def system_settings():
    config = SystemConfig.query.first()
    if not config:
        return jsonify({"error": "系统配置不存在"}), 404

    action = request.form.get("action", "")

    if action == "update_time_slots":
        visit_times = request.form.get("visit_times", "09:00-11:00,14:00-16:00").strip()
        config.visit_times = visit_times
        _write_log('operation', '系统设置', '预约时段设置', f'更新预约时段为 {visit_times}')

    elif action == "update_archive_types":
        archive_types = request.form.get("archive_types", "").strip()
        config.archive_types = archive_types
        _write_log('operation', '系统设置', '档案类型设置', f'更新档案类型')

    elif action == "update_campuses":
        campuses = request.form.get("campuses", "").strip()
        config.campuses = campuses
        _write_log('operation', '系统设置', '校区设置', f'更新校区列表')

    elif action == "update_daily_limit":
        daily_limit = request.form.get("daily_limit", 50, type=int)
        config.daily_limit = daily_limit
        _write_log('operation', '系统设置', '每日限额设置', f'更新每日限额为 {daily_limit}')

    db.session.commit()
    flash("系统设置已更新")
    return redirect(url_for("admin.dashboard", active_tab="settings"))
